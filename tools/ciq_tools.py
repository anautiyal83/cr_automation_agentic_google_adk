"""CIQ and failure classification tools."""
from __future__ import annotations

import json
import logging

logger = logging.getLogger(__name__)


def read_ciq_structure(ciq_file_path: str) -> dict:
    """Read CIQ Excel and return sheet names + headers + sample rows."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ciq_file_path, data_only=True, read_only=True)
        result = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=5, values_only=True))
            if not rows:
                continue
            headers = [str(h) if h else "" for h in rows[0]]
            samples = []
            for row in rows[1:4]:
                sample = {headers[i]: str(v) if v else "" for i, v in enumerate(row) if i < len(headers)}
                samples.append(sample)
            result[sheet_name] = {"headers": headers, "sample_rows": samples}
        wb.close()
        return result
    except Exception as e:
        return {"error": str(e)}


def classify_failure_type(failures: list[dict]) -> dict:
    """Classify execution failures into workflow vs simulator categories."""
    WORKFLOW_PATTERNS = [
        "cannot create property", "mapping values are not allowed",
        "patternsyntaxexception", "named capturing group",
        "node not found: null", "unable to find property",
    ]
    SIMULATOR_PATTERNS = [
        "connection refused", "ssh login failed", "command not found",
        "prompt not reached", "failed to start mockwire",
    ]

    workflow_issues, simulator_issues = [], []

    for f in failures:
        text = json.dumps(f, default=str).lower()
        is_wf = any(p in text for p in WORKFLOW_PATTERNS)
        is_sim = any(p in text for p in SIMULATOR_PATTERNS)

        if is_wf:
            workflow_issues.append(f.get("command", "unknown"))
        elif is_sim:
            simulator_issues.append(f.get("command", "unknown"))
        elif f.get("success") and f.get("validation_status") == "failed":
            workflow_issues.append(f.get("command", "unknown"))
        else:
            simulator_issues.append(f.get("command", "unknown"))

    classification = "simulator" if len(simulator_issues) > len(workflow_issues) else "workflow"
    return {
        "classification": classification,
        "workflow_issues": workflow_issues,
        "simulator_issues": simulator_issues,
        "retry_target": 8 if classification == "simulator" else 6,
    }
