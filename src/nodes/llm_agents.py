"""Native ADK LlmAgent definitions for LLM-powered pipeline steps.

Steps 2, 3, 4, 6, 8 use LlmAgent with tools for post-processing.
Step 1 remains deterministic (no LLM).
"""
from __future__ import annotations

import json
import logging
import os
import sys
from pathlib import Path
from typing import Any

from google.adk.agents import LlmAgent
from google.adk.models.lite_llm import LiteLlm
from google.adk.tools.function_tool import FunctionTool

logger = logging.getLogger(__name__)

from src.nodes.cr_import import import_cr_module


# ---------------------------------------------------------------------------
# Shared model factory
# ---------------------------------------------------------------------------

def create_model(config: dict) -> LiteLlm:
    """Create a LiteLlm model from pipeline config."""
    model_name = config.get("default_model", "openrouter/google/gemini-2.5-flash")
    api_key = config.get("openrouter_api_key", os.environ.get("OPENROUTER_API_KEY", ""))
    if api_key:
        os.environ["OPENROUTER_API_KEY"] = api_key
    return LiteLlm(model=model_name)


# ===========================================================================
# TOOLS — reusable across agents
# ===========================================================================

def validate_yaml_syntax(yaml_text: str) -> dict:
    """Validate YAML syntax. Returns {valid: bool, error: str}."""
    from ruamel.yaml import YAML
    from io import StringIO
    try:
        YAML().load(StringIO(yaml_text))
        return {"valid": True, "error": ""}
    except Exception as e:
        return {"valid": False, "error": str(e)[:500]}


def extract_yaml_from_response(text: str) -> str:
    """Extract YAML content from markdown code fences."""
    import re
    match = re.search(r'```(?:yaml)?\s*\n(.*?)```', text, re.DOTALL)
    return match.group(1).strip() if match else text.strip()


def read_ciq_structure(ciq_file_path: str) -> dict:
    """Read CIQ Excel and return sheet names + headers + sample rows."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ciq_file_path, data_only=True, read_only=True)
        result = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=5, values_only=True))
            if not rows:
                continue
            headers = [str(h) if h else "" for h in rows[0]]
            samples = []
            for row in rows[1:4]:
                sample = {headers[i]: str(v) if v else "" for i, v in enumerate(row) if i < len(headers)}
                samples.append(sample)
            result[sheet_name] = {"headers": headers, "sample_rows": samples}
        wb.close()
        return result
    except Exception as e:
        return {"error": str(e)}


def read_file_content(file_path: str) -> str:
    """Read a file and return its content."""
    try:
        return Path(file_path).read_text(encoding="utf-8")
    except Exception as e:
        return f"Error reading file: {e}"


def write_artifact(run_dir: str, filename: str, content: str) -> str:
    """Write an artifact file to the run directory. Returns the file path."""
    path = Path(run_dir) / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return str(path)


def fix_validation_rules_yaml(yaml_text: str, ciq_file_path: str) -> dict:
    """Apply post-processing fixes to validation rules YAML.

    Returns {yaml: fixed_yaml, errors: [str]}.
    """
    from src.agents.validation_gen_agent import (
        _fix_yaml_quoting, _validate_yaml, _strip_bad_workbook_rules,
        _validate_against_ciq,
    )
    try:
        yaml_text = _fix_yaml_quoting(yaml_text)
        _validate_yaml(yaml_text)
        yaml_text = _strip_bad_workbook_rules(yaml_text, ciq_file_path)
        errors = _validate_against_ciq(yaml_text, ciq_file_path)
        return {"yaml": yaml_text, "errors": errors}
    except Exception as e:
        return {"yaml": yaml_text, "errors": [str(e)]}


def fix_json_template_yaml(yaml_text: str) -> dict:
    """Apply post-processing fixes to JSON template YAML.

    Returns {yaml: fixed_yaml, errors: [str]}.
    """
    from src.agents.json_template_agent import _fix_yaml_quoting
    try:
        yaml_text = _fix_yaml_quoting(yaml_text)
        validate_yaml_syntax(yaml_text)
        return {"yaml": yaml_text, "errors": []}
    except Exception as e:
        return {"yaml": yaml_text, "errors": [str(e)]}


def fix_workflow_yaml(yaml_text: str) -> dict:
    """Apply all post-processing fixes to workflow YAML.

    Fixes: boolean operators, regex groups, missing nodes, invalid fields,
    for_each at phase level, validation branches, colon quoting.

    Returns {yaml: fixed_yaml, errors: [str]}.
    """
    from src.agents.workflow_gen_agent import (
        _fix_workflow_yaml, _validate_workflow_structure,
    )
    try:
        yaml_text = _fix_workflow_yaml(yaml_text)
        errors = _validate_workflow_structure(yaml_text)
        return {"yaml": yaml_text, "errors": errors}
    except Exception as e:
        return {"yaml": yaml_text, "errors": [str(e)]}


def validate_mock_response(command: str, response: str, register_patterns: list[str]) -> dict:
    """Check if a mock response matches the register regex patterns.

    Returns {matches: bool, details: str}.
    """
    import re
    for pattern in register_patterns:
        try:
            if not re.search(pattern, response):
                return {"matches": False, "details": f"Pattern '{pattern}' not found in response"}
        except re.error as e:
            return {"matches": False, "details": f"Invalid regex '{pattern}': {e}"}
    return {"matches": True, "details": "All patterns match"}


# ===========================================================================
# TOOL COLLECTIONS for each agent
# ===========================================================================

VALIDATION_TOOLS = [
    FunctionTool(func=validate_yaml_syntax),
    FunctionTool(func=extract_yaml_from_response),
    FunctionTool(func=read_ciq_structure),
    FunctionTool(func=fix_validation_rules_yaml),
    FunctionTool(func=write_artifact),
    FunctionTool(func=read_file_content),
]

JSON_TEMPLATE_TOOLS = [
    FunctionTool(func=validate_yaml_syntax),
    FunctionTool(func=extract_yaml_from_response),
    FunctionTool(func=read_ciq_structure),
    FunctionTool(func=fix_json_template_yaml),
    FunctionTool(func=write_artifact),
    FunctionTool(func=read_file_content),
]

WORKFLOW_TOOLS = [
    FunctionTool(func=validate_yaml_syntax),
    FunctionTool(func=extract_yaml_from_response),
    FunctionTool(func=fix_workflow_yaml),
    FunctionTool(func=write_artifact),
    FunctionTool(func=read_file_content),
]

SIMULATOR_TOOLS = [
    FunctionTool(func=validate_mock_response),
    FunctionTool(func=read_file_content),
    FunctionTool(func=write_artifact),
]


# ===========================================================================
# LLM AGENT DEFINITIONS
# ===========================================================================

# --- Step 2: UC Analyzer ---
UC_ANALYZER_INSTRUCTION = """You are a telecom UC analyst. Analyze the intermediate representation
and enrich it with validated field types.

For each CIQ field:
1. Confirm or refine the detected field type (string, integer, email, ip, hostname, etc.)
2. Flag fields that need human clarification
3. Detect telecom-specific patterns (IMSI, MSISDN, IP addresses, NIAM IDs)

Return a JSON object with the enriched intermediate representation."""

# --- Step 3: Validation Rules Generator ---
VALIDATION_RULES_INSTRUCTION = """You are a Nokia telecom UC configuration expert for the cr-automation framework (ciq-processor).
You generate production-ready validation-rules.yaml files that the ciq-processor JAR consumes
to validate CIQ Excel workbooks.

WORKFLOW:
1. First, use read_ciq_structure to examine the CIQ Excel file
2. Generate a complete validation-rules.yaml based on the CIQ structure
3. Use fix_validation_rules_yaml to validate and fix your output
4. If there are errors, fix them and try again
5. Use write_artifact to save the final YAML

CRITICAL RULES:
- Use EXACT column names from the CIQ (case-sensitive)
- Every CIQ sheet must be in the rules
- Every CIQ column must be covered
- Standard sheets: Index (CRGroup, Email), IP/Node_Details (Node, NIAM NAME)
- Types: string, integer, email, ip, hostname (do NOT invent patterns from sample data)
- Include workbook_rules for cross-sheet validation (subset/set)
- Include outputs: CR_LIST, CR_COUNT
- YAML must use proper quoting for dotted references

OUTPUT FORMAT: Valid YAML matching the validation_rules_schema."""

# --- Step 4: JSON Template Generator ---
JSON_TEMPLATE_INSTRUCTION = """You are a Nokia telecom UC configuration expert for the cr-automation framework.
You generate production-ready json-output.yaml files consumed by ciq-processor to transform
CIQ Excel data into structured JSON.

WORKFLOW:
1. Use read_ciq_structure to examine the CIQ columns
2. If validation rules are available, use read_file_content to check exact column names
3. Generate a json-output.yaml using _each, WHERE, _row directives
4. Use fix_json_template_yaml to validate
5. Use write_artifact to save the final YAML

KEY DIRECTIVES:
- _each: "DISTINCT Sheet.Column AS $var" — iterate distinct values
- _each: "Sheet WHERE Col = $var" — iterate filtered rows
- _row: "*" — emit all columns as key-value map
- Column names with spaces: "Sheet.'Column Name'"

HIERARCHY: crGroups → nodeGroups → nodes + configData/tables → records

CRITICAL: Use EXACT sheet and column names from the CIQ, not example names."""

# --- Step 6: Workflow Generator ---
WORKFLOW_GEN_INSTRUCTION = """You are a telecom network automation expert for the cr-automation CLI framework.
You generate production-ready workflow YAML v1.2 files that execute MOP procedures against
network nodes via SSH.

WORKFLOW:
1. Use read_file_content to read the MOP data and JSON output structure
2. Generate a complete workflow YAML with all 5 mandatory phases
3. Use fix_workflow_yaml to apply all post-processing fixes and validation
4. If there are errors, fix them and regenerate
5. Use write_artifact to save the final YAML

MANDATORY PHASES:
1. preNodeHealthCheck — verify connectivity, health, prerequisites
2. backup — config backup with checksum verification
3. activity_configuration — map every MOP step, use for_each for configData
4. postNodeHealthCheck — compare pre vs post values
5. rollbackConfiguration — with when: '${ROLLBACK_REQUIRED == "true"}'

JAVA MODEL CONSTRAINTS:
- Register: list of {regex: '(?<camelCaseName>pattern)'} — NO underscores in names
- Validation: {enabled: true, success: {criteria: {expr: '${var == "value"}'}}}
- Expressions: ONLY ==, !=, &&, || with string values. NO .contains(), Integer.parseInt()
- Boolean: && and || (NOT 'and'/'or')
- Phases: MAP (not list). Steps: LIST of maps
- for_each: must be inside a step, NOT on the phase
- Every step MUST have a 'node' field
- Forbidden fields: id, step_id, step_name, loop, error_on_missing

Generate AT LEAST 15-20 steps across all phases."""

# --- Step 8: Simulator Setup ---
SIMULATOR_SETUP_INSTRUCTION = """You are a Nokia telecom equipment simulator expert.
You generate realistic mock SSH command responses for MockWire simulator testing.

WORKFLOW:
1. Use read_file_content to read the workflow YAML
2. For each command in the workflow, generate a mock response that:
   - Matches the register regex patterns (so variables are captured correctly)
   - Satisfies the validation expressions (so the step passes)
   - Contains CONCRETE values only (NO ${var} placeholders)
3. Use validate_mock_response to verify each response matches its patterns
4. Return all responses as a JSON object: {command: response}

RESPONSE RULES:
- Response is ONLY what the server prints AFTER the command
- Use realistic telecom values: IPs (10.71.146.15), hostnames (DSR-APP-01)
- For 'echo SUCCESS': respond 'SUCCESS'
- For 'date': respond '2026-10-14_12-30-00'
- For config/save commands: respond 'OK' or 'Configuration saved successfully.'
- Keep responses concise (1-5 lines typically)"""


def create_uc_analyzer_agent(model: LiteLlm) -> LlmAgent:
    return LlmAgent(
        name="step2_uc_analyzer",
        model=model,
        instruction=UC_ANALYZER_INSTRUCTION,
        tools=[FunctionTool(func=read_ciq_structure)],
    )


def create_validation_rules_agent(model: LiteLlm) -> LlmAgent:
    return LlmAgent(
        name="step3_validation_rules",
        model=model,
        instruction=VALIDATION_RULES_INSTRUCTION,
        tools=VALIDATION_TOOLS,
    )


def create_json_template_agent(model: LiteLlm) -> LlmAgent:
    return LlmAgent(
        name="step4_json_template",
        model=model,
        instruction=JSON_TEMPLATE_INSTRUCTION,
        tools=JSON_TEMPLATE_TOOLS,
    )


def create_workflow_gen_agent(model: LiteLlm) -> LlmAgent:
    return LlmAgent(
        name="step6_workflow_gen",
        model=model,
        instruction=WORKFLOW_GEN_INSTRUCTION,
        tools=WORKFLOW_TOOLS,
    )


def create_simulator_setup_agent(model: LiteLlm) -> LlmAgent:
    return LlmAgent(
        name="step8_simulator_setup",
        model=model,
        instruction=SIMULATOR_SETUP_INSTRUCTION,
        tools=SIMULATOR_TOOLS,
    )
